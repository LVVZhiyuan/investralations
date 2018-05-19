# -*- coding: utf-8 -*-
"""
Spyder Editor

This is a temporary script file.
"""

import os
import win32com
from win32com.client import Dispatch, constants
from docx import Document
w = win32com.client.Dispatch('Word.Application')
def parse_doc(f):
    doc = w.Documents.Open( FileName = f )
    t = doc.Tables[0] # 根据文件中的图表选择信息
    name = t.Rows[0].Cells[1].Range.Text  
    situation = t.Rows[0].Cells[5].Range.Text
    people = t.Rows[1].Cells[1].Range.Text
    title = t.Rows[1].Cells[3].Range.Text  
    print name, situation, people,title
    doc.Close()
def parse_docx(f):

    d = Document(f)
    t = d.tables[0]
    name = t.cell(0,1).text
    situation = t.cell(0,8).text
    people = t.cell(1,2).text
    title = t.cell(1,8).text
    print name, situation, people,title
if __name__ == "__main__":
  
    PATH = "D:\wjh\" # 
    doc_files = os.listdir(PATH)
    for doc in doc_files:
        if os.path.splitext(doc)[1] == '.docx':
            try:
                parse_docx(PATH+'\\'+doc)
            except Exception as e:
                print e
        elif os.path.splitext(doc)[1] == '.doc':
            try:
                parse_doc(PATH+'\\'+doc)
            except Exception as e:
                print e